"""Compara los KPIs realizados de este año con los del año pasado -dos
archivos de ventas distintos, cargados por separado (éste y
"Comparativa interanual")-. La BBDD no trae columna de año, sólo el
nombre del mes (se repite cada año), así que la comparación empareja
por nombre de mes (ENERO de este año con ENERO del año pasado), no por
fecha real.

El año pasado NO lleva aplicadas las correcciones de "Ajustes
manuales": esa tabla no distingue año -está pensada para corregir el
año en curso-, así que reutilizar kpis_periodo() (que sí las aplica)
también para el año pasado mezclaría por error una corrección de este
año con los datos del año pasado."""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import metrics

FUENTE = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FUENTE, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FUENTE, bold=True, size=14)
BASE_FONT = Font(name=FUENTE)
PCT_FORMAT = "0.0%"
POS_FILL = PatternFill("solid", fgColor="E8F5E9")
NEG_FILL = PatternFill("solid", fgColor="FDEAEA")

METRICAS = metrics.METRICAS_CALCULADAS  # retail, bps, remarketing, bev, wholesale_uc, wholesale_yuc, ventas_totales

METRICAS_LEGIBLES = {
    "retail": "Retail",
    "bps": "BPS/MN",
    "remarketing": "Remarketing",
    "bev": "BEV",
    "wholesale_uc": "Wholesale UC",
    "wholesale_yuc": "Wholesale YUC",
    "ventas_totales": "Ventas totales",
}


def _totales_sin_ajustes(resumen: pd.DataFrame, codigo_dealer: int, marca: str, periodo: str, mes_referencia: str | None) -> dict:
    """Igual que metrics.kpis_periodo() pero sin aplicar ajustes_manuales -ver docstring del módulo."""
    meses = metrics.meses_de_periodo(periodo, mes_referencia)
    sub = resumen[
        (resumen["codigo_dealer"] == codigo_dealer)
        & (resumen["marca"] == marca)
        & (resumen["mes"].isin(meses))
    ]
    return {m: int(sub[m].sum()) for m in METRICAS}


def tabla_comparativa(resumen_actual: pd.DataFrame, resumen_anterior: pd.DataFrame, dealers: pd.DataFrame, marca: str, periodo: str, mes_referencia: str | None = None) -> pd.DataFrame:
    """Una fila por concesionario de `marca`, con el total de cada
    métrica este año (con ajustes manuales aplicados, igual que el
    resto de la app) y el año pasado (sin ajustes), más la variación
    absoluta y en %."""
    col = "vende_bmw" if marca == "BMW" else "vende_mini"
    filas = []
    for _, d in dealers[dealers[col] == "Si"].sort_values("concesionario").iterrows():
        codigo = int(d["codigo_dealer"])
        actual = metrics.kpis_periodo(resumen_actual, codigo, marca, periodo, mes_referencia)
        anterior = _totales_sin_ajustes(resumen_anterior, codigo, marca, periodo, mes_referencia)
        fila = {"codigo_dealer": codigo, "concesionario": d["concesionario"], "grupo_propietario": d["grupo_propietario"]}
        for m in METRICAS:
            val_actual, val_anterior = actual[m], anterior[m]
            fila[f"{m}_actual"] = val_actual
            fila[f"{m}_anterior"] = val_anterior
            fila[f"{m}_var_abs"] = val_actual - val_anterior
            fila[f"{m}_var_pct"] = ((val_actual - val_anterior) / val_anterior) if val_anterior else None
        filas.append(fila)
    return pd.DataFrame(filas)


def totales_comparativa(tabla: pd.DataFrame, metrica: str) -> dict:
    """Totales agregados de toda la tabla para una métrica -para las tarjetas resumen."""
    if tabla.empty:
        return {"actual": 0, "anterior": 0, "var_abs": 0, "var_pct": None}
    actual = int(tabla[f"{metrica}_actual"].sum())
    anterior = int(tabla[f"{metrica}_anterior"].sum())
    var_abs = actual - anterior
    var_pct = (var_abs / anterior) if anterior else None
    return {"actual": actual, "anterior": anterior, "var_abs": var_abs, "var_pct": var_pct}


def build_workbook(tabla: pd.DataFrame, marca: str, periodo: str, etiqueta_actual: str = "Este año", etiqueta_anterior: str = "Año pasado") -> bytes:
    """Excel descargable: una fila por concesionario, con
    Actual/Año pasado/Variación para cada métrica calculada."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparativa interanual"
    ws["A1"] = f"Comparativa interanual — {marca} — {periodo}"
    ws["A1"].font = TITLE_FONT

    header_row = 3
    headers = ["Código", "Concesionario", "Grupo propietario"]
    for m in METRICAS:
        legible = METRICAS_LEGIBLES[m]
        headers += [f"{legible} ({etiqueta_actual})", f"{legible} ({etiqueta_anterior})", f"{legible} Var.", f"{legible} Var. %"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    fila_excel = header_row + 1
    for _, r in tabla.iterrows():
        col = 1
        for v in (int(r["codigo_dealer"]), r["concesionario"], r["grupo_propietario"]):
            ws.cell(row=fila_excel, column=col, value=v).font = BASE_FONT
            col += 1
        for m in METRICAS:
            actual, anterior = int(r[f"{m}_actual"]), int(r[f"{m}_anterior"])
            var_abs, var_pct = int(r[f"{m}_var_abs"]), r[f"{m}_var_pct"]
            ws.cell(row=fila_excel, column=col, value=actual).font = BASE_FONT
            col += 1
            ws.cell(row=fila_excel, column=col, value=anterior).font = BASE_FONT
            col += 1
            c = ws.cell(row=fila_excel, column=col, value=var_abs)
            c.font = BASE_FONT
            if var_abs > 0:
                c.fill = POS_FILL
            elif var_abs < 0:
                c.fill = NEG_FILL
            col += 1
            c = ws.cell(row=fila_excel, column=col, value=var_pct)
            c.number_format = PCT_FORMAT
            c.font = BASE_FONT
            col += 1
        fila_excel += 1

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for j in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
